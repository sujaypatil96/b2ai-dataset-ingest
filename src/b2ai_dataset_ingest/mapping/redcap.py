"""REDCap export primitives -- the wide-form counterpart to :mod:`mapping.omop`.

AI-READI withholds four variables from its public releases (sex, race/ethnicity, medications,
5-digit zip -- docs.aireadi.org/docs/3/controlled-variables) and delivers them to approved
users not as OMOP tables but as raw REDCap exports: one Excel workbook per form, straight out
of the data-capture instrument. Nothing here is AI-READI-specific. REDCap's export conventions
are documented and stable, so another REDCap-sourced supplement reuses this module and writes
only its own config -- the same split :mod:`mapping.omop` makes for OMOP CDM.

Three conventions are encoded, each checked against the AI-READI export on 2026-10-06:

- **A checkbox field exports one column per choice**, ``<field>___<code>`` with the code
  LOWER-CASED (choice ``C41261`` becomes ``race___c41261``), each holding ``1`` or ``0``. The
  data dictionary spells the code as entered, so :func:`checkbox_column` lower-cases.
- **A repeating instrument exports one row per instance**, keyed by the record id plus
  ``redcap_repeat_instrument`` and ``redcap_repeat_instance``.
- **Excel types cells, and the types lie.** An integer code arrives as ``int`` on one row and
  ``str`` on the next; a dose typed as ``1-2`` arrives as a *datetime*, because Excel
  auto-converted it on entry. :func:`cell_text` reduces every cell to text so the readers see
  the same ``dict[str, str]`` a CSV export gives, rendering a datetime in a form no numeric
  parser accepts -- the mangled dose is then counted as unparseable instead of being read as
  the day of the month.

Reading ``.xlsx`` needs ``openpyxl`` (the ``excel`` extra); a ``.csv``/``.tsv`` export of the
same sheet needs nothing. Nothing here logs a cell value.
"""

from __future__ import annotations

import csv
import re
from collections.abc import Iterator
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

#: Export formats a protected-supplement config may point at.
SUPPORTED_SUFFIXES = frozenset({".xlsx", ".xlsm", ".csv", ".tsv"})

#: REDCap's separator between a checkbox field and its choice code in export column names.
CHECKBOX_SEP = "___"

#: Spellings REDCap and Excel use for a ticked checkbox.
_CHECKED = frozenset({"1", "1.0", "true", "checked"})

_FLOAT_ZERO = re.compile(r"^(\d+)\.0+$")
_DATETIME_TEXT = re.compile(r"^\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?$")


class ExcelSupportMissing(RuntimeError):
    """``openpyxl`` is not installed, so an ``.xlsx`` export cannot be read."""


def locate_files(folder: Path, pattern: str) -> list[Path]:
    """Files under ``folder`` matching the config's glob, in sorted order.

    Excel's ``~$`` lock files are skipped -- one is left behind whenever a workbook is open --
    and so is anything outside :data:`SUPPORTED_SUFFIXES`, so a pattern can name a form
    without committing to an extension (``*Medications*Protected*``) and both the ``.xlsx``
    delivered and a ``.csv`` conversion of it match. The caller decides what to do with zero
    or several hits; both are configuration errors it should report rather than guess past.
    """
    if not pattern:
        return []
    return sorted(
        path
        for path in Path(folder).glob(pattern)
        if path.is_file()
        and not path.name.startswith("~$")
        and path.suffix.lower() in SUPPORTED_SUFFIXES
    )


def checkbox_column(field: str, code: Any) -> str:
    """The export column for one checkbox choice: ``race`` + ``C41261`` -> ``race___c41261``."""
    return f"{field}{CHECKBOX_SEP}{str(code).strip().lower()}"


def is_checked(value: Any) -> bool:
    """True if a checkbox export cell says the choice was ticked."""
    return cell_text(value).lower() in _CHECKED


def normalize_id(value: Any) -> str:
    """A record id as the OMOP tables spell it: ``1001``, ``"1001"`` and ``1001.0`` -> ``"1001"``.

    Excel types an id column as int, or as float once a blank appears in it, while the OMOP
    CSVs carry the same id as text. Joining on anything but the normalized text would silently
    drop every row of the supplement.
    """
    text = cell_text(value)
    match = _FLOAT_ZERO.match(text)
    return match.group(1) if match else text


def cell_text(value: Any) -> str:
    """Reduce a typed cell to the text a CSV export would carry; blank for ``None``."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime | date | time):
        # Deliberately not a number: a dose Excel turned into a date must fail numeric parsing
        # and be counted, never be read as the day of the month.
        return value.isoformat()
    return str(value).strip()


def looks_like_datetime(value: Any) -> bool:
    """True for a datetime-typed cell, or for text shaped like one (a CSV conversion of it)."""
    if isinstance(value, datetime | date | time):
        return True
    return bool(_DATETIME_TEXT.match(cell_text(value)))


def read_header(path: Path, sheet: int | str | None = None) -> list[str]:
    """The export's column names, in order; empty for an empty sheet."""
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        worksheet, workbook = _worksheet(path, sheet)
        try:
            for cells in worksheet.iter_rows(values_only=True):
                return [cell_text(cell) for cell in cells]
            return []
        finally:
            workbook.close()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"unsupported export format {path.suffix!r}; expected .xlsx/.csv/.tsv")
    with open(path, newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh, delimiter="\t" if suffix == ".tsv" else ",")
        return [cell.strip() for cell in next(reader, [])]


def stream_rows(
    path: Path, sheet: int | str | None = None, typed: bool = False
) -> Iterator[dict[str, Any]]:
    """Yield one ``{column: cell}`` dict per non-blank data row.

    ``typed=False`` (the readers) gives every cell as text via :func:`cell_text`;
    ``typed=True`` (the validator) gives the raw cell so Excel's auto-conversions can be
    counted. Rows blank in every column are skipped: Excel exports routinely carry trailing
    ones.
    """
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        yield from _stream_xlsx(path, sheet, typed)
    elif suffix in {".csv", ".tsv"}:
        yield from _stream_delimited(path, "\t" if suffix == ".tsv" else ",", typed)
    else:
        raise ValueError(f"unsupported export format {path.suffix!r}; expected .xlsx/.csv/.tsv")


def _worksheet(path: Path, sheet: int | str | None):
    try:
        import openpyxl
    except ImportError as exc:
        raise ExcelSupportMissing(
            f"{path.name} is an Excel workbook and openpyxl is not installed; install the "
            "'excel' extra (uv sync --extra excel) or export the sheet to CSV"
        ) from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet is None:
        return workbook.worksheets[0], workbook
    if isinstance(sheet, int):
        return workbook.worksheets[sheet], workbook
    return workbook[sheet], workbook


def _stream_xlsx(path: Path, sheet: int | str | None, typed: bool) -> Iterator[dict[str, Any]]:
    worksheet, workbook = _worksheet(path, sheet)
    try:
        rows = worksheet.iter_rows(values_only=True)
        header_cells = next(rows, None)
        if header_cells is None:
            return
        header = [cell_text(cell) for cell in header_cells]
        for cells in rows:
            if all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in cells):
                continue
            values = list(cells) + [None] * (len(header) - len(cells))
            yield {
                name: (value if typed else cell_text(value))
                for name, value in zip(header, values, strict=False)
                if name
            }
    finally:
        workbook.close()


def _stream_delimited(path: Path, delimiter: str, typed: bool) -> Iterator[dict[str, Any]]:
    with open(path, newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh, delimiter=delimiter):
            if all((value or "").strip() == "" for value in row.values() if value is not None):
                continue
            yield {
                name.strip(): (value if typed else cell_text(value))
                for name, value in row.items()
                if name
            }


__all__ = [
    "CHECKBOX_SEP",
    "SUPPORTED_SUFFIXES",
    "ExcelSupportMissing",
    "cell_text",
    "checkbox_column",
    "is_checked",
    "locate_files",
    "looks_like_datetime",
    "normalize_id",
    "read_header",
    "stream_rows",
]
